"""Parse CSV, Excel, PDF, Word, TXT and image metadata into DataFrames."""

from __future__ import annotations

import io
import json
from typing import BinaryIO

import pandas as pd
from fastapi import HTTPException

ALLOWED_EXTENSIONS = {
    ".csv",
    ".tsv",
    ".xlsx",
    ".xls",
    ".pdf",
    ".txt",
    ".md",
    ".json",
    ".doc",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
    ".bmp",
}


def _ext(filename: str) -> str:
    name = (filename or "").lower()
    for ext in sorted(ALLOWED_EXTENSIONS, key=len, reverse=True):
        if name.endswith(ext):
            return ext
    return ""


def parse_upload(filename: str, file_bytes: bytes) -> pd.DataFrame:
    ext = _ext(filename)
    if not ext:
        raise HTTPException(
            status_code=400,
            detail={
                "fr": "Format non supporté. Utilisez CSV, Excel, PDF, Word, TXT, JSON ou une image.",
                "en": "Unsupported format. Use CSV, Excel, PDF, Word, TXT, JSON or an image.",
            },
        )

    buffer = io.BytesIO(file_bytes)

    try:
        if ext in {".csv", ".tsv"}:
            return _read_csv(buffer, sep="\t" if ext == ".tsv" else None)
        if ext in {".xlsx", ".xls"}:
            return _read_excel(buffer, ext)
        if ext == ".pdf":
            return _read_pdf(file_bytes)
        if ext in {".txt", ".md"}:
            return _read_text(file_bytes, filename)
        if ext == ".json":
            return _read_json(file_bytes)
        if ext in {".docx", ".doc"}:
            return _read_word(file_bytes, ext)
        if ext in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}:
            return _read_image_meta(filename, file_bytes)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail={
                "fr": f"Impossible de lire le fichier: {error}",
                "en": f"Unable to read the file: {error}",
            },
        ) from error

    raise HTTPException(status_code=400, detail="Unsupported file")


def _read_csv(buffer: BinaryIO, sep: str | None = None) -> pd.DataFrame:
    seps = [sep] if sep else [",", ";", "\t", "|"]
    for candidate in seps:
        buffer.seek(0)
        try:
            df = pd.read_csv(buffer, sep=candidate, engine="python")
            if df.shape[1] >= 1:
                return _normalize_columns(df)
        except Exception:
            continue
    buffer.seek(0)
    df = pd.read_csv(buffer)
    return _normalize_columns(df)


def _read_excel(buffer: BinaryIO, ext: str) -> pd.DataFrame:
    buffer.seek(0)
    engine = "openpyxl" if ext == ".xlsx" else "xlrd"
    df = pd.read_excel(buffer, engine=engine)
    return _normalize_columns(df)


def _read_pdf(file_bytes: bytes) -> pd.DataFrame:
    try:
        import pdfplumber
    except ImportError as error:
        raise HTTPException(
            status_code=500,
            detail={
                "fr": "pdfplumber n'est pas installé.",
                "en": "pdfplumber is not installed.",
            },
        ) from error

    frames: list[pd.DataFrame] = []
    text_chunks: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table) < 2:
                    continue
                header = [
                    str(cell).strip() if cell is not None else f"col_{i}"
                    for i, cell in enumerate(table[0])
                ]
                rows = table[1:]
                frame = pd.DataFrame(rows, columns=header)
                frames.append(frame)
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_chunks.append(page_text.strip())

    if frames:
        df = pd.concat(frames, ignore_index=True)
        return _normalize_columns(df)

    if text_chunks:
        return _normalize_columns(
            pd.DataFrame({"contenu": text_chunks})
        )

    raise HTTPException(
        status_code=400,
        detail={
            "fr": "Aucun tableau ni texte détecté dans le PDF.",
            "en": "No table or text detected in the PDF.",
        },
    )


def _read_text(file_bytes: bytes, filename: str) -> pd.DataFrame:
    text = file_bytes.decode("utf-8", errors="replace")
    # Try CSV-like text first
    try:
        df = pd.read_csv(io.StringIO(text))
        if df.shape[1] > 1:
            return _normalize_columns(df)
    except Exception:
        pass
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise HTTPException(
            status_code=400,
            detail={
                "fr": "Le fichier texte est vide.",
                "en": "The text file is empty.",
            },
        )
    return _normalize_columns(
        pd.DataFrame({"fichier": [filename], "contenu": ["\n".join(lines[:200])]})
    )


def _read_json(file_bytes: bytes) -> pd.DataFrame:
    payload = json.loads(file_bytes.decode("utf-8", errors="replace"))
    if isinstance(payload, list):
        df = pd.json_normalize(payload)
    elif isinstance(payload, dict):
        df = pd.json_normalize(payload)
    else:
        df = pd.DataFrame({"valeur": [payload]})
    return _normalize_columns(df)


def _read_word(file_bytes: bytes, ext: str) -> pd.DataFrame:
    if ext == ".doc":
        raise HTTPException(
            status_code=400,
            detail={
                "fr": "Les fichiers .doc anciens ne sont pas supportés. Convertissez en .docx.",
                "en": "Legacy .doc files are not supported. Please convert to .docx.",
            },
        )
    try:
        from docx import Document
    except ImportError as error:
        raise HTTPException(
            status_code=500,
            detail={
                "fr": "python-docx n'est pas installé.",
                "en": "python-docx is not installed.",
            },
        ) from error

    document = Document(io.BytesIO(file_bytes))
    frames: list[pd.DataFrame] = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if len(rows) >= 2:
            header = rows[0]
            frames.append(pd.DataFrame(rows[1:], columns=header))

    if frames:
        return _normalize_columns(pd.concat(frames, ignore_index=True))

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if not paragraphs:
        raise HTTPException(
            status_code=400,
            detail={
                "fr": "Le document Word ne contient aucune donnée exploitable.",
                "en": "The Word document contains no usable data.",
            },
        )
    return _normalize_columns(pd.DataFrame({"contenu": paragraphs}))


def _read_image_meta(filename: str, file_bytes: bytes) -> pd.DataFrame:
    return _normalize_columns(
        pd.DataFrame(
            [
                {
                    "type": "image",
                    "filename": filename,
                    "size_bytes": len(file_bytes),
                    "note": (
                        "Image jointe. Posez une question sur ce que représente "
                        "l'image ou les données qu'elle contient."
                    ),
                }
            ]
        )
    )


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    cleaned = df.copy()
    cleaned.columns = [
        str(col).strip() if str(col).strip() else f"colonne_{i + 1}"
        for i, col in enumerate(cleaned.columns)
    ]
    cleaned = cleaned.dropna(how="all").reset_index(drop=True)
    if cleaned.empty:
        raise HTTPException(
            status_code=400,
            detail={
                "fr": "Le fichier ne contient aucune donnée exploitable.",
                "en": "The file contains no usable data.",
            },
        )
    return cleaned
